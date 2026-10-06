from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, session
import os
import csv
import io
import re
import psycopg
from psycopg.rows import dict_row
from psycopg.errors import UniqueViolation
from datetime import datetime
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash

BASE=os.path.dirname(os.path.abspath(__file__))
UPLOADS=os.path.join(BASE,"uploads")
os.makedirs(UPLOADS,exist_ok=True)

app=Flask(__name__)
app.secret_key=os.environ.get("FILETRACK_SECRET","change-this-secret-key")

LGAS=["Katsina A","Katsina B","Batagarawa","Kaita","Jibia","Rimi","Charanchi"]
BRANCHES=["CIM","CW&HS","CDS"]
PRIORITIES=["Normal","High","Urgent"]
ROLES=["Administrator","Zonal Inspector","LGI Officer","Branch Official","Supporting Staff"]

def db():
    database_url=os.environ.get("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL environment variable is not configured.")

    # Render PostgreSQL external connections use SSL. Add this only when
    # the supplied DATABASE_URL does not already specify an sslmode.
    if "sslmode=" not in database_url.lower():
        separator = "&" if "?" in database_url else "?"
        database_url += separator + "sslmode=require"

    last_error = None
    for attempt in range(3):
        try:
            return psycopg.connect(
                database_url,
                row_factory=dict_row,
                connect_timeout=10,
                keepalives=1,
                keepalives_idle=30,
                keepalives_interval=10,
                keepalives_count=3,
            )
        except psycopg.OperationalError as exc:
            last_error = exc
            if attempt < 2:
                import time
                time.sleep(1)

    raise last_error



@app.after_request
def integrate_inspection_ui(response):
    """Ensure clean response headers without injecting UI cards."""
    try:
        if not session.get("user_id") or "text/html" not in (response.content_type or ""):
            return response
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    except Exception:
        pass
    return response

def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def init_db():
    """Create/upgrade the PostgreSQL schema used by FileTrack."""
    with db() as c:
        c.execute("""
        CREATE TABLE IF NOT EXISTS files(
          id BIGSERIAL PRIMARY KEY,
          file_id TEXT UNIQUE NOT NULL,
          title TEXT NOT NULL,
          reference_no TEXT,
          lga TEXT NOT NULL,
          lgi_name TEXT,
          received_by TEXT NOT NULL,
          received_at TEXT NOT NULL,
          priority TEXT DEFAULT 'Normal',
          description TEXT,
          attachment TEXT,
          status TEXT DEFAULT 'Received',
          current_location TEXT DEFAULT 'Katsina Zonal Office',
          created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS movements(
          id BIGSERIAL PRIMARY KEY,
          file_id TEXT NOT NULL,
          from_location TEXT NOT NULL,
          to_location TEXT NOT NULL,
          forwarded_by TEXT,
          receiving_officer TEXT,
          action TEXT,
          forwarded_at TEXT NOT NULL,
          acknowledged_at TEXT,
          status TEXT DEFAULT 'Forwarded',
          remarks TEXT,
          submitted_by TEXT,
          submitted_at TEXT
        );

        CREATE TABLE IF NOT EXISTS audit_logs(
          id BIGSERIAL PRIMARY KEY,
          file_id TEXT NOT NULL,
          action TEXT NOT NULL,
          official TEXT,
          details TEXT,
          created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS users(
          id BIGSERIAL PRIMARY KEY,
          full_name TEXT NOT NULL,
          username TEXT UNIQUE NOT NULL,
          password_hash TEXT NOT NULL,
          role TEXT NOT NULL,
          office TEXT,
          lga TEXT,
          active INTEGER DEFAULT 1,
          created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS inspections(
          id BIGSERIAL PRIMARY KEY,
          inspection_id TEXT UNIQUE NOT NULL,
          lga TEXT NOT NULL,
          ppa_employer TEXT,
          corps_member TEXT,
          subject TEXT NOT NULL,
          findings TEXT NOT NULL,
          recommendations TEXT,
          inspected_by TEXT NOT NULL,
          inspected_at TEXT NOT NULL,
          status TEXT DEFAULT 'Submitted to ZI'
        );

        CREATE TABLE IF NOT EXISTS reports(
          id BIGSERIAL PRIMARY KEY,
          report_id TEXT UNIQUE NOT NULL,
          lga TEXT NOT NULL,
          report_type TEXT NOT NULL,
          ppa_employer TEXT,
          corps_member TEXT,
          subject TEXT NOT NULL,
          report_body TEXT NOT NULL,
          issued_by TEXT NOT NULL,
          issued_at TEXT NOT NULL,
          status TEXT DEFAULT 'Issued'
        );

        CREATE TABLE IF NOT EXISTS ppa_establishments(
          id BIGSERIAL PRIMARY KEY,
          name TEXT NOT NULL,
          lga TEXT NOT NULL,
          active INTEGER DEFAULT 1,
          created_by TEXT,
          created_at TEXT NOT NULL,
          UNIQUE(name,lga)
        );

        CREATE TABLE IF NOT EXISTS corps_members(
          id BIGSERIAL PRIMARY KEY,
          state_code TEXT UNIQUE NOT NULL,
          full_name TEXT NOT NULL,
          gender TEXT,
          discipline TEXT,
          ppa_id BIGINT REFERENCES ppa_establishments(id) ON DELETE SET NULL,
          batch TEXT,
          stream TEXT,
          phone TEXT,
          status TEXT DEFAULT 'Active',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS inspection_assignments(
          id BIGSERIAL PRIMARY KEY,
          ppa_id BIGINT NOT NULL REFERENCES ppa_establishments(id) ON DELETE CASCADE,
          supporting_staff_id BIGINT NOT NULL REFERENCES users(id) ON DELETE RESTRICT,
          assigned_by BIGINT NOT NULL REFERENCES users(id) ON DELETE RESTRICT,
          assigned_at TEXT NOT NULL,
          status TEXT DEFAULT 'Active'
        );

        CREATE UNIQUE INDEX IF NOT EXISTS ux_active_inspection_ppa
          ON inspection_assignments(ppa_id) WHERE status='Active';

        CREATE TABLE IF NOT EXISTS inspection_records(
          id BIGSERIAL PRIMARY KEY,
          inspection_id TEXT UNIQUE NOT NULL,
          corps_member_id BIGINT NOT NULL REFERENCES corps_members(id) ON DELETE RESTRICT,
          ppa_id BIGINT NOT NULL REFERENCES ppa_establishments(id) ON DELETE RESTRICT,
          inspected_by BIGINT NOT NULL REFERENCES users(id) ON DELETE RESTRICT,
          attendance_status TEXT NOT NULL,
          query_issued INTEGER DEFAULT 0,
          other_reason TEXT,
          remarks TEXT,
          inspected_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS corps_member_next_of_kin(
          id BIGSERIAL PRIMARY KEY,
          state_code TEXT UNIQUE NOT NULL REFERENCES corps_members(state_code) ON DELETE CASCADE,
          lga TEXT NOT NULL,
          next_of_kin_name TEXT NOT NULL,
          relationship TEXT,
          phone TEXT,
          address TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        );

        ALTER TABLE movements ADD COLUMN IF NOT EXISTS submitted_by TEXT;
        ALTER TABLE movements ADD COLUMN IF NOT EXISTS submitted_at TEXT;
        """)

def user_count():
    c=db()
    n=c.execute("SELECT COUNT(*) FROM users").fetchone()["count"]
    c.close()
    return n

def current_user():
    uid=session.get("user_id")
    if not uid:
        return None
    c=db()
    u=c.execute("SELECT * FROM users WHERE id=%s AND active=1",(uid,)).fetchone()
    c.close()
    return u

@app.context_processor
def inject_user():
    return {"current_user": current_user()}

def login_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user():
            return redirect(url_for("login",next=request.path))
        return view(*args,**kwargs)
    return wrapped

def admin_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        u=current_user()
        if not u:
            return redirect(url_for("login",next=request.path))
        if u["role"]!="Administrator":
            flash("Administrator access is required.","error")
            return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped

def official_name():
    u=current_user()
    return u["full_name"] if u else ""

def role_is(role):
    u=current_user()
    return bool(u and u["role"] == role)

def can_receive_register():
    u=current_user()
    return bool(u and u["role"] in ("Administrator", "Zonal Inspector", "Supporting Staff"))

def can_operate_file_movement():
    u=current_user()
    return bool(u and u["role"] in ("Administrator", "Zonal Inspector"))

def supporting_staff_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user():
            return redirect(url_for("login",next=request.path))
        if current_user()["role"] != "Supporting Staff":
            flash("This action is reserved for Supporting Staff responsible for physical file delivery.","error")
            return redirect(url_for("detail",file_id=kwargs.get("file_id","")))
        return view(*args,**kwargs)
    return wrapped


def lgi_report_only(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user(): return redirect(url_for("login",next=request.path))
        if current_user()["role"] == "LGI Officer":
            flash("LGI Officers have report-only access. Reports are issued by the Zonal Inspector for their assigned LGA.","error")
            return redirect(url_for("reports"))
        return view(*args,**kwargs)
    return wrapped

def registry_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user():
            return redirect(url_for("login",next=request.path))
        if not can_receive_register():
            flash("Only the Administrator, Zonal Inspector or Supporting Staff may register incoming LGI files.","error")
            return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped

def movement_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user():
            return redirect(url_for("login",next=request.path))
        if not can_operate_file_movement():
            flash("Only the Administrator or Zonal Inspector may forward or return files. Supporting Staff register incoming files and confirm physical delivery.","error")
            return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped

def next_file_id():
    c=db()
    n=c.execute("SELECT COUNT(*) AS n FROM files").fetchone()["n"]
    c.close()
    return f"KZO-FM-{datetime.now().year}-{n+1:04d}"

def next_inspection_id():
    c=db(); n=c.execute("SELECT COUNT(*) FROM inspections").fetchone()["count"]; c.close(); return f"KZO-INSP-{datetime.now().year}-{n+1:04d}"

def next_report_id():
    c=db(); n=c.execute("SELECT COUNT(*) FROM reports").fetchone()["count"]; c.close(); return f"KZO-RPT-{datetime.now().year}-{n+1:04d}"

def inspection_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user(): return redirect(url_for("login",next=request.path))
        if current_user()["role"] != "Supporting Staff":
            flash("Inspection is assigned to Supporting Staff. LGI Officers receive Zonal reports instead.","error"); return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped

def report_management_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not current_user(): return redirect(url_for("login",next=request.path))
        if current_user()["role"] not in ("Administrator", "Zonal Inspector"):
            flash("Only the Administrator or Zonal Inspector may issue reports to LGI Officers.","error"); return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped

@app.route("/health")
def health():
    try:
        with db() as c:
            row = c.execute(
                "SELECT current_database() AS database_name, current_user AS database_user"
            ).fetchone()
        return {
            "status": "ok",
            "database": row["database_name"],
            "user": row["database_user"],
        }
    except Exception as exc:
        return {"status": "error", "message": str(exc)}, 503


@app.route("/login",methods=["GET","POST"])
def login():
    if current_user():
        return redirect(url_for("dashboard"))
    if request.method=="POST":
        username=request.form.get("username","").strip()
        password=request.form.get("password","")
        c=db()
        u=c.execute("SELECT * FROM users WHERE username=%s",(username,)).fetchone()
        c.close()
        if u and u["active"] and check_password_hash(u["password_hash"],password):
            session.clear()
            session["user_id"]=u["id"]
            session["username"]=u["username"]
            session["full_name"]=u["full_name"]
            session["role"]=u["role"]
            nxt=request.args.get("next") or request.form.get("next") or url_for("dashboard")
            return redirect(nxt)
        flash("Invalid username or password, or the account is inactive.","error")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.","success")
    return redirect(url_for("login"))

@app.route("/setup",methods=["GET","POST"])
def setup():
    if user_count()>0:
        flash("Initial setup has already been completed.","error")
        return redirect(url_for("login"))
    if request.method=="POST":
        full_name=request.form.get("full_name","").strip()
        username=request.form.get("username","").strip()
        password=request.form.get("password","")
        confirm=request.form.get("confirm_password","")
        if not full_name or not username or not password:
            flash("All required fields must be completed.","error")
        elif len(password)<8:
            flash("Password must be at least 8 characters.","error")
        elif password!=confirm:
            flash("Passwords do not match.","error")
        else:
            c=db()
            c.execute("""INSERT INTO users
                (full_name,username,password_hash,role,office,lga,active,created_at)
                VALUES(%s,%s,%s,%s,%s,%s,1,%s)""",
                (full_name,username,generate_password_hash(password),
                 "Administrator","Katsina Zonal Office","",now()))
            c.commit()
            c.close()
            flash("Administrator account created. Please log in.","success")
            return redirect(url_for("login"))
    return render_template("setup.html")

@app.route("/users",methods=["GET","POST"])
@admin_required
def users():
    if request.method=="POST":
        full_name=request.form.get("full_name","").strip()
        username=request.form.get("username","").strip()
        password=request.form.get("password","")
        role=request.form.get("role","").strip()
        office=request.form.get("office","").strip()
        lga=request.form.get("lga","").strip()
        if not full_name or not username or not password or role not in ROLES:
            flash("Full name, username, password and valid role are required.","error")
        elif len(password)<8:
            flash("Password must be at least 8 characters.","error")
        else:
            c=db()
            try:
                c.execute("""INSERT INTO users
                    (full_name,username,password_hash,role,office,lga,active,created_at)
                    VALUES(%s,%s,%s,%s,%s,%s,1,%s)""",
                    (full_name,username,generate_password_hash(password),
                     role,office,lga,now()))
                c.commit()
                flash(f"{full_name} account created successfully.","success")
            except UniqueViolation:
                flash("That username already exists.","error")
            finally:
                c.close()
        return redirect(url_for("users"))
    c=db()
    rows=c.execute("SELECT * FROM users ORDER BY full_name").fetchall()
    c.close()
    return render_template("users.html",users=rows,roles=ROLES,lgas=LGAS)

@app.route("/users/<int:user_id>/toggle",methods=["POST"])
@admin_required
def toggle_user(user_id):
    if user_id==current_user()["id"]:
        flash("You cannot deactivate your own account.","error")
        return redirect(url_for("users"))
    c=db()
    c.execute("UPDATE users SET active=CASE WHEN active=1 THEN 0 ELSE 1 END WHERE id=%s",(user_id,))
    c.commit()
    c.close()
    flash("Account status updated.","success")
    return redirect(url_for("users"))

@app.route("/")
@login_required
def dashboard():
    u=current_user(); c=db()
    if u["role"] == "LGI Officer":
        reports=c.execute("SELECT * FROM reports WHERE lga=%s ORDER BY id DESC LIMIT 20",(u["lga"],)).fetchall(); c.close()
        return render_template("dashboard.html",lgi_reports=reports)
    total=c.execute("SELECT COUNT(*) FROM files").fetchone()["count"]; received=c.execute("SELECT COUNT(*) FROM files WHERE status='Received'").fetchone()["count"]; forwarded=c.execute("SELECT COUNT(*) FROM files WHERE status='Forwarded'").fetchone()["count"]; submitted=c.execute("SELECT COUNT(*) FROM files WHERE status='Submitted'").fetchone()["count"]; acknowledged=c.execute("SELECT COUNT(*) FROM files WHERE status='Acknowledged'").fetchone()["count"]; returned=c.execute("SELECT COUNT(*) FROM files WHERE status='Returned'").fetchone()["count"]
    stats={"total":total,"received":received,"forwarded":forwarded,"submitted":submitted,"ack":acknowledged,"returned":returned,"Total":total,"Received":received,"Forwarded":forwarded,"Submitted":submitted,"Acknowledged":acknowledged,"Returned":returned}
    branch_counts={b:c.execute("SELECT COUNT(*) FROM files WHERE current_location=%s",(b,)).fetchone()["count"] for b in BRANCHES}; branches=[(b,branch_counts[b]) for b in BRANCHES]; recent=c.execute("SELECT * FROM files ORDER BY id DESC LIMIT 8").fetchall(); inspection_count=c.execute("SELECT COUNT(*) FROM inspections").fetchone()["count"]; pending_inspections=c.execute("SELECT COUNT(*) FROM inspections WHERE status='Submitted to ZI'").fetchone()["count"]; report_count=c.execute("SELECT COUNT(*) FROM reports").fetchone()["count"]
    my_inspections=0
    if u["role"] == "Supporting Staff":
        my_inspections=c.execute("SELECT COUNT(*) FROM inspection_assignments WHERE supporting_staff_id=%s AND status='Active'",(u["id"],)).fetchone()["count"]
    c.close()
    return render_template("dashboard.html",stats=stats,branches=branches,branch_counts=branch_counts,recent=recent,inspection_count=inspection_count,pending_inspections=pending_inspections,report_count=report_count,my_inspections=my_inspections)

@app.route("/files")
@login_required
@lgi_report_only
def files_page():
    if current_user()["role"] == "Supporting Staff":
        return redirect(url_for("supporting_file_search"))

    q=request.args.get("q","").strip()
    status=request.args.get("status","").strip()
    branch=request.args.get("branch","").strip()
    c=db()
    sql="SELECT * FROM files WHERE 1=1"; args=[]
    if q:
        sql += """ AND (file_id LIKE %s OR title LIKE %s OR reference_no LIKE %s
                    OR lga LIKE %s OR lgi_name LIKE %s OR current_location LIKE %s)"""
        args += [f"%{q}%"]*6
    if status:
        sql+=" AND status=%s"; args.append(status)
    if branch:
        sql+=" AND current_location=%s"; args.append(branch)
    sql+=" ORDER BY id DESC"
    rows=c.execute(sql,args).fetchall()
    c.close()
    return render_template("files.html",rows=rows,q=q,status=status,
                           branch=branch,branches=BRANCHES)

@app.route("/lgi-send-file",methods=["GET","POST"])
@login_required
def lgi_send_file():
    """Allow an LGI to submit a file to the Zonal Office for their assigned LGA only."""
    u=current_user()
    if not u or u["role"] != "LGI Officer":
        flash("This submission page is reserved for LGI Officers.","error")
        return redirect(url_for("dashboard"))
    assigned_lga=(u.get("lga") or "").strip()
    if assigned_lga not in LGAS:
        flash("Your account does not have a valid assigned LGA. Contact the Administrator.","error")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        title=request.form.get("title","").strip()
        if not title:
            flash("File/document title is required.","error")
            return redirect(url_for("lgi_send_file"))
        t=now(); fid=next_file_id(); attachment=None
        uploaded=request.files.get("attachment")
        if uploaded and uploaded.filename:
            attachment=f"{fid}_{os.path.basename(uploaded.filename)}"
            uploaded.save(os.path.join(UPLOADS,attachment))
        c=db()
        try:
            c.execute("""INSERT INTO files
                (file_id,title,reference_no,lga,lgi_name,received_by,received_at,priority,
                 description,attachment,status,current_location,created_at)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (fid,title,request.form.get("reference_no"),assigned_lga,
                 u["full_name"],u["full_name"],t,request.form.get("priority","Normal"),
                 request.form.get("remarks") or request.form.get("description"),attachment,
                 "Submitted","Katsina Zonal Office",t))
            c.execute("""INSERT INTO movements
                (file_id,from_location,to_location,forwarded_by,receiving_officer,action,
                 forwarded_at,status,remarks,submitted_by,submitted_at)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (fid,assigned_lga,"Katsina Zonal Office",u["full_name"],"Zonal Inspector",
                 "LGI FILE SUBMISSION",t,"Submitted",request.form.get("remarks"),u["full_name"],t))
            c.execute("""INSERT INTO audit_logs
                (file_id,action,official,details,created_at) VALUES(%s,%s,%s,%s,%s)""",
                (fid,"FILE SUBMITTED BY LGI",u["full_name"],
                 f"Submitted from assigned LGA: {assigned_lga} to Katsina Zonal Office",t))
            c.commit()
        except Exception:
            c.rollback()
            raise
        finally:
            c.close()
        flash(f"{fid} submitted to the Zonal Office successfully.","success")
        return redirect(url_for("dashboard"))

    return render_template("receive.html",lgas=[assigned_lga],priorities=PRIORITIES,
                           lgi_scope=True,assigned_lga=assigned_lga)



@app.route("/lgi-corps-search")
@login_required
def lgi_corps_search():
    """LGI directory restricted server-side to the officer's assigned LGA."""
    u=current_user()
    if u["role"] != "LGI Officer":
        flash("The corps member directory is available to LGI Officers only.","error")
        return redirect(url_for("dashboard"))
    assigned_lga=(u.get("lga") or "").strip()
    if assigned_lga not in LGAS:
        flash("Your account has no valid assigned LGA. Contact the Administrator.","error")
        return redirect(url_for("dashboard"))
    q=request.args.get("q","").strip()
    rows=[]
    if q:
        c=db()
        like=f"%{q}%"
        rows=c.execute("""
            SELECT cm.state_code,cm.full_name,cm.gender,cm.discipline,cm.batch,cm.stream,
                   cm.phone,cm.status,p.name AS ppa_name,p.lga
            FROM corps_members cm
            JOIN ppa_establishments p ON p.id=cm.ppa_id
            WHERE p.lga=%s AND COALESCE(cm.status,'Active') ILIKE 'Active'
              AND (cm.state_code ILIKE %s OR cm.full_name ILIKE %s OR COALESCE(cm.phone,'') ILIKE %s
                   OR p.name ILIKE %s OR COALESCE(cm.discipline,'') ILIKE %s)
            ORDER BY cm.full_name LIMIT 250
        """,(assigned_lga,like,like,like,like,like)).fetchall()
        c.close()
    return render_template("lgi_corps_search.html",assigned_lga=assigned_lga,q=q,rows=rows)


@app.route("/lgi-corps/<path:state_code>")
@login_required
def lgi_corps_detail(state_code):
    """Show an individual corps record only when it belongs to the LGI's LGA."""
    u=current_user()
    if u["role"] != "LGI Officer":
        flash("The corps member directory is available to LGI Officers only.","error")
        return redirect(url_for("dashboard"))
    assigned_lga=(u.get("lga") or "").strip()
    if assigned_lga not in LGAS:
        flash("Your account has no valid assigned LGA. Contact the Administrator.","error")
        return redirect(url_for("dashboard"))
    c=db()
    row=c.execute("""
        SELECT cm.state_code,cm.full_name,cm.gender,cm.discipline,cm.batch,cm.stream,
               cm.phone,cm.status,cm.created_at,cm.updated_at,p.name AS ppa_name,p.lga
        FROM corps_members cm JOIN ppa_establishments p ON p.id=cm.ppa_id
        WHERE cm.state_code=%s AND p.lga=%s
    """,(state_code,assigned_lga)).fetchone()
    c.close()
    if not row:
        return "Corps member not found in your assigned LGA.",404
    return render_template("lgi_corps_detail.html",r=row,assigned_lga=assigned_lga)


@app.route("/lgi-file-status")
@login_required
def lgi_file_status():
    """Track only files submitted by the logged-in LGI to the Zonal Office."""
    u=current_user()
    if u["role"] != "LGI Officer":
        flash("File submission tracking is available to LGI Officers only.","error")
        return redirect(url_for("dashboard"))
    assigned_lga=(u.get("lga") or "").strip()
    if assigned_lga not in LGAS:
        flash("Your account has no valid assigned LGA. Contact the Administrator.","error")
        return redirect(url_for("dashboard"))
    q=request.args.get("q","").strip()
    c=db()
    params=[assigned_lga,u["full_name"]]
    sql="""
        SELECT f.file_id,f.title,f.reference_no,f.lga,f.status,f.current_location,
               f.received_at,f.created_at,f.priority,f.description,
               m.action AS latest_action,m.status AS latest_movement_status,
               m.to_location AS latest_to_location,m.forwarded_at AS latest_movement_at,
               m.remarks AS latest_remarks
        FROM files f
        LEFT JOIN LATERAL (
            SELECT action,status,to_location,forwarded_at,remarks
            FROM movements WHERE movements.file_id=f.file_id ORDER BY id DESC LIMIT 1
        ) m ON TRUE
        WHERE f.lga=%s AND f.lgi_name=%s
    """
    # Bind scope parameters in SQL order; the first two are used in the scope clause below.
    params=[]
    if q:
        like=f"%{q}%"
        sql += " AND (f.file_id ILIKE %s OR f.title ILIKE %s OR COALESCE(f.reference_no,'') ILIKE %s OR f.status ILIKE %s)"
        params.extend([like,like,like,like])
    sql += " ORDER BY f.id DESC LIMIT 500"
    rows=c.execute(sql,(assigned_lga,u["full_name"],*params)).fetchall()
    c.close()
    return render_template("lgi_file_status.html",rows=rows,q=q,assigned_lga=assigned_lga)

@app.route("/receive",methods=["GET","POST"])
@registry_required
def receive():
    if request.method=="POST":
        title=request.form.get("title","").strip()
        if not title:
            flash("File/document title is required.","error")
            return redirect(url_for("receive"))
        t=now(); fid=next_file_id(); attachment=None
        uploaded=request.files.get("attachment")
        if uploaded and uploaded.filename:
            attachment=f"{fid}_{os.path.basename(uploaded.filename)}"
            uploaded.save(os.path.join(UPLOADS,attachment))
        official=official_name()
        c=db()
        c.execute("""INSERT INTO files
        (file_id,title,reference_no,lga,lgi_name,received_by,received_at,priority,
         description,attachment,status,current_location,created_at)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
        (fid,title,request.form.get("reference_no"),request.form["lga"],
         request.form.get("lgi_name"),official,t,request.form.get("priority","Normal"),
         request.form.get("description"),attachment,"Received",
         "Katsina Zonal Office",t))
        c.execute("""INSERT INTO audit_logs
        (file_id,action,official,details,created_at) VALUES(%s,%s,%s,%s,%s)""",
        (fid,"FILE RECEIVED",official,"Received from "+request.form["lga"],t))
        c.commit(); c.close()
        flash(fid+" registered successfully.","success")

        if current_user()["role"] == "Supporting Staff":
            return redirect(
                url_for("supporting_file_view", file_id=fid)
            )

        return redirect(url_for("detail",file_id=fid))
    return render_template("receive.html",lgas=LGAS,priorities=PRIORITIES)

@app.route("/supporting-file-search")
@login_required
def supporting_file_search():
    u = current_user()

    if u["role"] != "Supporting Staff":
        flash("This search is available to Supporting Staff only.","error")
        return redirect(url_for("dashboard"))

    q = request.args.get("q","").strip()
    rows = []

    if q:
        c = db()
        rows = c.execute("""
            SELECT file_id,title,reference_no,lga,priority,status,
                   description,attachment
            FROM files
            WHERE file_id ILIKE %s
               OR reference_no ILIKE %s
            ORDER BY id DESC
        """,(f"%{q}%",f"%{q}%")).fetchall()
        c.close()

    return render_template(
        "supporting_file_search.html",
        rows=rows,
        q=q
    )


@app.route("/supporting-file/<file_id>")
@login_required
def supporting_file_view(file_id):
    u = current_user()

    if u["role"] != "Supporting Staff":
        flash("This page is available to Supporting Staff only.","error")
        return redirect(url_for("dashboard"))

    c = db()
    f = c.execute("""
        SELECT file_id,title,reference_no,lga,priority,status,
               description,attachment
        FROM files
        WHERE file_id=%s
    """,(file_id,)).fetchone()
    c.close()

    if not f:
        flash("File not found.","error")
        return redirect(url_for("supporting_file_search"))

    return render_template(
        "supporting_file_view.html",
        f=f
    )


@app.route("/file/<file_id>")
@login_required
@lgi_report_only
def detail(file_id):
    if current_user()["role"] == "Supporting Staff":
        return redirect(
            url_for("supporting_file_view", file_id=file_id)
        )

    c=db()
    f=c.execute("SELECT * FROM files WHERE file_id=%s",(file_id,)).fetchone()
    if not f:
        c.close(); return "File not found",404
    movements=c.execute("SELECT * FROM movements WHERE file_id=%s ORDER BY id DESC",(file_id,)).fetchall()
    audit=c.execute("SELECT * FROM audit_logs WHERE file_id=%s ORDER BY id DESC",(file_id,)).fetchall()
    c.close()
    return render_template("detail.html",f=f,movements=movements,audit=audit,branches=BRANCHES)

@app.route("/forward/<file_id>",methods=["POST"])
@movement_required
def forward(file_id):
    t=now(); c=db()
    f=c.execute("SELECT * FROM files WHERE file_id=%s",(file_id,)).fetchone()
    if not f:
        c.close(); return "File not found",404
    official=official_name(); to=request.form["to_location"]
    c.execute("""INSERT INTO movements
    (file_id,from_location,to_location,forwarded_by,receiving_officer,
     action,forwarded_at,status,remarks) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
    (file_id,f["current_location"],to,official,request.form.get("receiving_officer"),
     request.form.get("action"),t,"Forwarded",request.form.get("remarks")))
    c.execute("UPDATE files SET current_location=%s,status='Forwarded' WHERE file_id=%s",(to,file_id))
    c.execute("""INSERT INTO audit_logs
    (file_id,action,official,details,created_at) VALUES(%s,%s,%s,%s,%s)""",
    (file_id,"FILE FORWARDED",official,f"{f['current_location']} → {to}",t))
    c.commit(); c.close()
    flash(f"{file_id} forwarded to {to}.","success")
    return redirect(url_for("detail",file_id=file_id))

@app.route("/ack/<file_id>",methods=["POST"])
@supporting_staff_required
def acknowledge(file_id):
    t=now(); official=official_name(); c=db()
    m=c.execute("""SELECT * FROM movements WHERE file_id=%s AND status='Forwarded'
                   ORDER BY id DESC LIMIT 1""",(file_id,)).fetchone()
    if not m:
        c.close(); flash("No pending forwarded file is awaiting physical-submission confirmation.","error")
        return redirect(url_for("supporting_file_view", file_id=file_id))
    c.execute("""UPDATE movements SET status='Submitted', submitted_at=%s, submitted_by=%s,
                 acknowledged_at=%s WHERE id=%s""",
              (t,official,t,m["id"]))
    c.execute("UPDATE files SET status='Submitted' WHERE file_id=%s",(file_id,))
    c.execute("""INSERT INTO audit_logs
    (file_id,action,official,details,created_at) VALUES(%s,%s,%s,%s,%s)""",
    (file_id,"PHYSICAL FILE SUBMITTED",official,
     f"Physical file submitted to {m['to_location']}",t))
    c.commit(); c.close()
    flash(
        f"Physical submission of {file_id} to {m['to_location']} confirmed.",
        "success"
    )

    return redirect(
        url_for("supporting_file_view", file_id=file_id)
    )

@app.route("/return/<file_id>",methods=["POST"])
@movement_required
def return_file(file_id):
    t=now(); official=official_name(); c=db()
    f=c.execute("SELECT * FROM files WHERE file_id=%s",(file_id,)).fetchone()
    if not f:
        c.close(); return "File not found",404
    c.execute("""UPDATE files SET current_location='Katsina Zonal Office',
                 status='Returned' WHERE file_id=%s""",(file_id,))
    c.execute("""INSERT INTO movements
    (file_id,from_location,to_location,forwarded_by,action,forwarded_at,status,remarks)
    VALUES(%s,%s,%s,%s,%s,%s,%s,%s)""",
    (file_id,f["current_location"],"Katsina Zonal Office",official,
     "Returned",t,"Returned",request.form.get("remarks")))
    c.execute("""INSERT INTO audit_logs
    (file_id,action,official,details,created_at) VALUES(%s,%s,%s,%s,%s)""",
    (file_id,"FILE RETURNED",official,
     request.form.get("remarks") or "Returned to Zonal Office",t))
    c.commit(); c.close()
    flash("File returned to Zonal Office.","success")
    return redirect(url_for("detail",file_id=file_id))

def inspection_management_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        u=current_user()
        if not u:
            return redirect(url_for("login",next=request.path))
        if u["role"] not in ("Administrator","Zonal Inspector"):
            flash("Only the Administrator or Zonal Inspector may manage inspection assignments.","error")
            return redirect(url_for("dashboard"))
        return view(*args,**kwargs)
    return wrapped


def _norm_header(value):
    return "_".join(str(value or "").strip().lower().replace("/"," ").replace("-"," ").split())


def _read_corps_upload(upload):
    filename=(upload.filename or "").lower()
    raw=upload.read()
    if filename.endswith(".csv"):
        text=raw.decode("utf-8-sig",errors="replace")
        return [{_norm_header(k): (v or "").strip() for k,v in row.items()} for row in csv.DictReader(io.StringIO(text))]
    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Excel upload requires openpyxl. Add openpyxl to requirements.txt and redeploy.")
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        ws=(wb["Expected"] if "Expected" in wb.sheetnames else wb["All Corps Members"] if "All Corps Members" in wb.sheetnames else wb.active)
        values=list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers=[_norm_header(x) for x in values[0]]
        rows=[]
        for vals in values[1:]:
            row={headers[i]: ("" if i>=len(vals) or vals[i] is None else str(vals[i]).strip())
                 for i in range(len(headers)) if headers[i]}
            if any(row.values()):
                rows.append(row)
        return rows
    raise RuntimeError("Unsupported file type. Please upload CSV or XLSX.")


def _infer_lga_from_filename(filename):
    base=os.path.splitext(os.path.basename(filename or ""))[0].strip().lower()
    base=re.sub(r"[_\-\s]*(?:a|b|c)$","",base)
    mapping={
        "batagarawa":"Batagarawa",
        "rimi":"Rimi",
        "kaita":"Kaita",
        "jibia":"Jibia",
        "charanchi":"Charanchi",
    }
    return mapping.get(base,"")


def _normalize_inspection_lga(value, component=""):
    v=str(value or "").strip().upper()
    comp=str(component or "").strip().upper()
    if v in ("KATSINA", "KATSINA MAIN", "KATSINA METROPOLIS"):
        if comp in ("KATSINA A", "KATSINA B"):
            return comp.title()
        return ""
    mapping={
        "BATAGARAWA":"Batagarawa", "RIMI":"Rimi", "KAITA":"Kaita",
        "JIBIA":"Jibia", "CHARANCHI":"Charanchi",
        "KATSINA A":"Katsina A", "KATSINA B":"Katsina B"
    }
    return mapping.get(v, str(value or "").strip())


def _legacy_corps_fields(row, filename="", selected_lga=""):
    state=_pick(row,"state_code","statecode","state code","state-code","code")

    name=_pick(row,"full_name","corps_member_name","corps member name","name","corps name")
    if not name:
        surname=_pick(row,"surname","last name","lastname")
        othernames=_pick(row,"othernames","other names","first name","firstname")
        name=" ".join(x for x in (surname,othernames) if x).strip()

    ppa=_pick(row,"ppa_establishment","ppa establishment","ppa / establishment",
              "ppa","ppa_employer","place_of_posting","place of posting",
              "employer","company","company name","establishment",
              "place of primary assignment")

    component=_pick(row,"component","zone component")
    raw_lga=_pick(row,"lga","local government","local government area","posted_lga","posted lga")
    lga=_normalize_inspection_lga(raw_lga,component)
    if not lga:
        fallback=(selected_lga or _infer_lga_from_filename(filename)).strip()
        lga=_normalize_inspection_lga(fallback,component) or fallback

    batch=_pick(row,"batch","batch year")
    if not batch and state:
        m=re.search(r"/(\d{2}[ABC])/?",state.upper())
        if m:
            batch=m.group(1)
    if batch:
        b=batch.strip().replace("_"," ")
        if re.fullmatch(r"[ABC]",b.upper()):
            batch=f"Batch_{b.upper()}"
        elif re.fullmatch(r"BATCH\s*[ABC]",b.upper()):
            batch="Batch_"+b[-1].upper()
        else:
            batch=batch.strip()

    return {
        "state":state,
        "name":name,
        "ppa":ppa,
        "lga":lga,
        "gender":_pick(row,"gender","sex"),
        "discipline":_pick(row,"discipline","course","qualification"),
        "batch":batch,
        "stream":_pick(row,"stream"),
        "phone":_pick(row,"phone","phone number","mobile","gsmno","gsm no","gsm no.","gsm"),
        "status":_pick(row,"status") or "Active",
    }


def _pick(row,*names):
    for name in names:
        key=_norm_header(name)
        if row.get(key): return str(row[key]).strip()
    return ""


def next_new_inspection_id():
    c=db(); n=c.execute("SELECT COUNT(*) FROM inspection_records").fetchone()["count"]; c.close()
    return f"KZO-INSP-{datetime.now().year}-{n+1:04d}"


@app.route("/inspections")
@login_required
def inspections():
    u=current_user()
    if u["role"] in ("Administrator","Zonal Inspector"):
        c=db()
        ppas=c.execute("""SELECT p.*, COUNT(cm.id) AS corps_count,
                    (SELECT COUNT(*) FROM inspection_assignments ia WHERE ia.ppa_id=p.id AND ia.status='Active') AS assigned
                    FROM ppa_establishments p LEFT JOIN corps_members cm ON cm.ppa_id=p.id
                    WHERE p.active=1 GROUP BY p.id ORDER BY p.lga,p.name""").fetchall()
        assignments=c.execute("""SELECT ia.*,p.name AS ppa_name,p.lga,u.full_name AS staff_name
                    FROM inspection_assignments ia JOIN ppa_establishments p ON p.id=ia.ppa_id
                    JOIN users u ON u.id=ia.supporting_staff_id
                    WHERE ia.status='Active' ORDER BY p.lga,p.name""").fetchall()
        staff=c.execute("SELECT id,full_name,username FROM users WHERE role='Supporting Staff' AND active=1 ORDER BY full_name").fetchall()
        recent=c.execute("""SELECT ir.*,cm.state_code,cm.full_name,p.name AS ppa_name,p.lga,u.full_name AS inspector_name
                    FROM inspection_records ir JOIN corps_members cm ON cm.id=ir.corps_member_id
                    JOIN ppa_establishments p ON p.id=ir.ppa_id JOIN users u ON u.id=ir.inspected_by
                    ORDER BY ir.id DESC LIMIT 30""").fetchall()
        c.close()
        return render_template("inspections.html",mode="management",ppas=ppas,assignments=assignments,staff=staff,recent=recent)
    if u["role"] == "Supporting Staff":
        c=db()
        assigned=c.execute("""SELECT p.id,p.name,p.lga,COUNT(cm.id) AS corps_count
                    FROM inspection_assignments ia JOIN ppa_establishments p ON p.id=ia.ppa_id
                    LEFT JOIN corps_members cm ON cm.ppa_id=p.id AND cm.status='Active'
                    WHERE ia.supporting_staff_id=%s AND ia.status='Active' AND p.active=1
                    GROUP BY p.id ORDER BY p.lga,p.name""",(u["id"],)).fetchall()
        mine=c.execute("""SELECT ir.*,cm.state_code,cm.full_name,p.name AS ppa_name,p.lga
                    FROM inspection_records ir JOIN corps_members cm ON cm.id=ir.corps_member_id
                    JOIN ppa_establishments p ON p.id=ir.ppa_id
                    WHERE ir.inspected_by=%s ORDER BY ir.id DESC LIMIT 50""",(u["id"],)).fetchall()
        c.close()
        return render_template("inspections.html",mode="staff",assigned=assigned,mine=mine)
    flash("Inspection access is not available to this role.","error")
    return redirect(url_for("dashboard"))


def _next_of_kin_fields(row, selected_lga=""):
    state=_pick(row,"state_code","statecode","state code","state-code","code")
    name=_pick(row,"next_of_kin_name","next of kin name","next_of_kin","next of kin","nok name","nok")
    relationship=_pick(row,"relationship","relationship to corps member","relation","next of kin relationship")
    phone=_pick(row,"next_of_kin_phone","next of kin phone","nok phone","nok_phone","phone","gsm","gsmno","phone number")
    address=_pick(row,"next_of_kin_address","next of kin address","nok address","address","residential address")
    raw_lga=_pick(row,"lga","local government","local government area","next of kin lga","nok lga")
    lga=_normalize_inspection_lga(raw_lga) or selected_lga
    return {"state":state.strip(),"name":name.strip(),"relationship":relationship.strip(),"phone":phone.strip(),"address":address.strip(),"lga":lga.strip()}


@app.route("/directory")
@login_required
def directory():
    u=current_user()
    if u["role"] == "LGI Officer":
        return render_template("directory.html", mode="lgi")
    if u["role"] in ("Administrator", "Zonal Inspector"):
        return render_template("directory.html", mode="zi")
    flash("The general Corps Members Directory is not available to Supporting Staff.","error")
    return redirect(url_for("dashboard"))


@app.route("/next-of-kin-search")
@login_required
def next_of_kin_search():
    u=current_user()
    if u["role"] != "LGI Officer":
        flash("Next of Kin search is available to LGI Officers only.","error")
        return redirect(url_for("dashboard"))
    assigned_lga=(u.get("lga") or "").strip()
    if assigned_lga not in LGAS:
        flash("Your account has no valid assigned LGA. Contact the Administrator.","error")
        return redirect(url_for("dashboard"))
    q=request.args.get("q","").strip()
    rows=[]
    if q:
        c=db(); like=f"%{q}%"
        rows=c.execute("""
            SELECT nok.state_code,nok.next_of_kin_name,nok.relationship,nok.phone,nok.address,nok.lga,
                   cm.full_name AS corps_member_name
            FROM corps_member_next_of_kin nok
            JOIN corps_members cm ON cm.state_code=nok.state_code
            JOIN ppa_establishments p ON p.id=cm.ppa_id
            WHERE p.lga=%s AND nok.lga=%s AND nok.state_code ILIKE %s
            ORDER BY cm.full_name LIMIT 250
        """,(assigned_lga,assigned_lga,like)).fetchall(); c.close()
    return render_template("next_of_kin_search.html",assigned_lga=assigned_lga,q=q,rows=rows)


@app.route("/next-of-kin/import",methods=["GET","POST"])
@inspection_management_required
def next_of_kin_import():
    if request.method == "POST":
        uploads=[u for u in request.files.getlist("nok_file") if u and u.filename]
        selected_lga=(request.form.get("lga") or "").strip()
        if not uploads:
            flash("Please select at least one CSV or XLSX file.","error"); return redirect(url_for("next_of_kin_import"))
        if selected_lga not in LGAS:
            flash("Please select the LGA for this Next of Kin import.","error"); return redirect(url_for("next_of_kin_import"))
        valid=[]; errors=0; processed=0; error_rows=[]
        try:
            for upload in uploads:
                filename=upload.filename
                try: rows=_read_corps_upload(upload)
                except Exception as exc: errors+=1; error_rows.append(f"{filename}: {exc}"); continue
                for idx,row in enumerate(rows,start=2):
                    f=_next_of_kin_fields(row,selected_lga); processed+=1
                    missing=[]
                    if not f["state"]: missing.append("State Code")
                    if not f["name"]: missing.append("Next of Kin Name")
                    if missing: errors+=1; error_rows.append(f"{filename} row {idx}: {', '.join(missing)} required"); continue
                    if f["lga"] != selected_lga: errors+=1; error_rows.append(f"{filename} row {idx}: LGA '{f['lga']}' does not match selected LGA '{selected_lga}'"); continue
                    valid.append(f)
            deduped={f["state"]:f for f in valid}; valid=list(deduped.values())
            if not valid:
                flash(f"Import completed: {processed} rows processed, 0 new, 0 updated, {errors} errors.","error")
                if error_rows: flash(" | ".join(error_rows[:8]),"error")
                return redirect(url_for("next_of_kin_import"))
            c=db(); t=now(); states=[f["state"] for f in valid]
            existing=c.execute("SELECT state_code FROM corps_members WHERE state_code = ANY(%s)",(states,)).fetchall()
            existing_states={r["state_code"] for r in existing}
            missing_states=[st for st in states if st not in existing_states]
            if missing_states:
                errors+=len(missing_states); error_rows.extend([f"{st}: Corps Member State Code does not exist in the directory" for st in missing_states[:8]])
                valid=[f for f in valid if f["state"] in existing_states]
            if valid:
                check=c.execute("""SELECT cm.state_code,p.lga FROM corps_members cm JOIN ppa_establishments p ON p.id=cm.ppa_id WHERE cm.state_code = ANY(%s)""",([f["state"] for f in valid],)).fetchall()
                lga_map={r["state_code"]:r["lga"] for r in check}; safe=[]
                for f in valid:
                    if lga_map.get(f["state"]) != selected_lga: errors+=1; error_rows.append(f"{f['state']}: Corps Member is not assigned to {selected_lga}")
                    else: safe.append(f)
                valid=safe
            if not valid:
                c.rollback(); c.close(); flash(f"Import completed: {processed} rows processed, 0 new, 0 updated, {errors} errors.","error")
                if error_rows: flash(" | ".join(error_rows[:8]),"error")
                return redirect(url_for("next_of_kin_import"))
            states=[f["state"] for f in valid]
            found=c.execute("SELECT state_code FROM corps_member_next_of_kin WHERE state_code = ANY(%s)",(states,)).fetchall(); existing_nok={r["state_code"] for r in found}
            new_count=sum(1 for f in valid if f["state"] not in existing_nok); updated_count=sum(1 for f in valid if f["state"] in existing_nok)
            values=[]; params=[]
            for f in valid:
                values.append("(%s,%s,%s,%s,%s,%s,%s,%s)"); params.extend([f["state"],f["lga"],f["name"],f["relationship"],f["phone"],f["address"],t,t])
            c.execute(f"""INSERT INTO corps_member_next_of_kin(state_code,lga,next_of_kin_name,relationship,phone,address,created_at,updated_at)
                VALUES {','.join(values)} ON CONFLICT (state_code) DO UPDATE SET lga=EXCLUDED.lga,next_of_kin_name=EXCLUDED.next_of_kin_name,relationship=EXCLUDED.relationship,phone=EXCLUDED.phone,address=EXCLUDED.address,updated_at=EXCLUDED.updated_at""",params)
            c.commit(); c.close()
            flash(f"Next of Kin import completed: {processed} rows processed, {new_count} new, {updated_count} updated, {errors} errors.","success" if errors==0 else "error")
            if error_rows: flash(" | ".join(error_rows[:8]),"error")
        except Exception as exc:
            try: c.rollback(); c.close()
            except Exception: pass
            flash(f"Next of Kin import failed: {exc}","error")
        return redirect(url_for("next_of_kin_import"))
    return render_template("next_of_kin_import.html",lgas=LGAS)


@app.route("/corps-members/import", methods=["GET", "POST"] )
@inspection_management_required
def corps_members_import():
    """Dedicated general corps-member directory import, separate from inspection navigation."""
    return _corps_import_handler(directory_mode=True)


@app.route("/inspection/import",methods=["GET","POST"])
@inspection_management_required
def inspection_import():
    return _corps_import_handler(directory_mode=False)


def _corps_import_handler(directory_mode=False):
    if request.method == "POST":
        uploads=[u for u in request.files.getlist("corps_file") if u and u.filename]
        selected_lga=(request.form.get("lga") or "").strip()
        if not uploads:
            flash("Please select at least one CSV or XLSX file.","error")
            return redirect(url_for("corps_members_import" if directory_mode else "inspection_import"))
        if selected_lga and selected_lga not in LGAS:
            flash("Please select a valid LGA.","error")
            return redirect(url_for("corps_members_import" if directory_mode else "inspection_import"))

        try:
            valid=[]
            errors=0
            processed=0
            error_rows=[]

            for upload in uploads:
                filename=upload.filename
                try:
                    rows=_read_corps_upload(upload)
                except Exception as exc:
                    errors += 1
                    error_rows.append(f"{filename}: {exc}")
                    continue

                for idx,row in enumerate(rows,start=2):
                    fields=_legacy_corps_fields(row,filename,selected_lga)
                    processed += 1
                    state,name,ppa,lga=(fields["state"],fields["name"],fields["ppa"],fields["lga"])
                    missing=[]
                    if not state: missing.append("State Code")
                    if not name: missing.append("Name")
                    if not ppa: missing.append("PPA/Establishment")
                    if not lga: missing.append("LGA")
                    if missing:
                        errors += 1
                        error_rows.append(f"{filename} row {idx}: {', '.join(missing)} required")
                        continue
                    if lga not in LGAS:
                        errors += 1
                        error_rows.append(f"{filename} row {idx}: Invalid LGA '{lga}'")
                        continue
                    valid.append(fields)

            if not valid:
                flash(f"Import completed: {processed} rows processed, 0 new, 0 updated, {errors} errors.","error")
                if error_rows:
                    flash(" | ".join(error_rows[:8]),"error")
                return redirect(url_for("corps_members_import" if directory_mode else "inspection_import"))

            deduped={}
            for f in valid:
                deduped[f["state"].strip()]=f
            valid=list(deduped.values())

            c=db()
            t=now()
            user_id=current_user()["id"]

            ppa_keys=[]
            seen_ppas=set()
            for f in valid:
                key=(f["ppa"].strip(),f["lga"].strip())
                if key not in seen_ppas:
                    seen_ppas.add(key)
                    ppa_keys.append(key)

            if ppa_keys:
                placeholders=",".join(["(%s,%s,1,%s,%s)"]*len(ppa_keys))
                params=[]
                for name,lga in ppa_keys:
                    params.extend([name,lga,user_id,t])
                c.execute(f"""
                    INSERT INTO ppa_establishments(name,lga,active,created_by,created_at)
                    VALUES {placeholders}
                    ON CONFLICT (name,lga) DO NOTHING
                """,params)

            ppa_where=",".join(["(%s,%s)"]*len(ppa_keys))
            ppa_params=[]
            for name,lga in ppa_keys:
                ppa_params.extend([name,lga])
            ppa_rows=c.execute(f"""
                SELECT id,name,lga FROM ppa_establishments
                WHERE (name,lga) IN ({ppa_where})
            """,ppa_params).fetchall()
            ppa_map={(r["name"].lower(),r["lga"].lower()):r["id"] for r in ppa_rows}

            states=[]
            seen_states=set()
            for f in valid:
                st=f["state"].strip()
                if st not in seen_states:
                    seen_states.add(st)
                    states.append(st)

            existing_states=set()
            for i in range(0,len(states),500):
                part=states[i:i+500]
                ph=",".join(["%s"]*len(part))
                found=c.execute(f"SELECT state_code FROM corps_members WHERE state_code IN ({ph})",part).fetchall()
                existing_states.update(r["state_code"] for r in found)

            new_count=sum(1 for f in valid if f["state"] not in existing_states)
            updated_count=sum(1 for f in valid if f["state"] in existing_states)

            columns=("state_code,full_name,gender,discipline,ppa_id,batch,stream,phone,status,created_at,updated_at")
            for offset in range(0,len(valid),500):
                chunk=valid[offset:offset+500]
                values=[]
                params=[]
                for f in chunk:
                    ppa_id=ppa_map.get((f["ppa"].strip().lower(),f["lga"].strip().lower()))
                    if not ppa_id:
                        raise RuntimeError(f"PPA could not be resolved: {f['ppa']} ({f['lga']})")
                    values.append("(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)")
                    params.extend([
                        f["state"],f["name"],f["gender"],f["discipline"],ppa_id,
                        f["batch"],f["stream"],f["phone"],f["status"],t,t
                    ])
                c.execute(f"""
                    INSERT INTO corps_members({columns})
                    VALUES {','.join(values)}
                    ON CONFLICT (state_code) DO UPDATE SET
                        full_name=EXCLUDED.full_name,
                        gender=EXCLUDED.gender,
                        discipline=EXCLUDED.discipline,
                        ppa_id=EXCLUDED.ppa_id,
                        batch=EXCLUDED.batch,
                        stream=EXCLUDED.stream,
                        phone=EXCLUDED.phone,
                        status=EXCLUDED.status,
                        updated_at=EXCLUDED.updated_at
                """,params)

            c.commit()
            c.close()

            msg=f"Import completed: {processed} rows processed, {new_count} new, {updated_count} updated, {errors} errors."
            flash(msg,"success" if errors==0 else "error")
            if error_rows:
                flash(" | ".join(error_rows[:8]),"error")

        except Exception as exc:
            try:
                c.rollback()
                c.close()
            except Exception:
                pass
            flash(f"Import failed: {exc}","error")
        return redirect(url_for("corps_members_import" if directory_mode else "inspection_import"))
    return render_template("corps_members_import.html" if directory_mode else "inspection_import.html", lgas=LGAS)


@app.route("/inspection/assign",methods=["POST"])
@inspection_management_required
def inspection_assign():
    ppa_id=request.form.get("ppa_id",type=int); staff_id=request.form.get("staff_id",type=int)
    if not ppa_id or not staff_id:
        flash("PPA/Establishment and Supporting Staff are required.","error"); return redirect(url_for("inspections"))
    c=db(); p=c.execute("SELECT * FROM ppa_establishments WHERE id=%s AND active=1",(ppa_id,)).fetchone(); staff=c.execute("SELECT * FROM users WHERE id=%s AND role='Supporting Staff' AND active=1",(staff_id,)).fetchone()
    if not p or not staff:
        c.close(); flash("Invalid PPA or Supporting Staff selected.","error"); return redirect(url_for("inspections"))
    c.execute("UPDATE inspection_assignments SET status='Reassigned' WHERE ppa_id=%s AND status='Active'",(ppa_id,))
    c.execute("""INSERT INTO inspection_assignments(ppa_id,supporting_staff_id,assigned_by,assigned_at,status)
                 VALUES(%s,%s,%s,%s,'Active')""",(ppa_id,staff_id,current_user()["id"],now()))
    c.commit(); c.close(); flash(f"{p['name']} assigned to {staff['full_name']}.","success"); return redirect(url_for("inspections"))


@app.route("/inspection/assignment/<int:assignment_id>/remove",methods=["POST"])
@inspection_management_required
def inspection_remove_assignment(assignment_id):
    c=db(); c.execute("UPDATE inspection_assignments SET status='Removed' WHERE id=%s",(assignment_id,)); c.commit(); c.close(); flash("Inspection assignment removed.","success"); return redirect(url_for("inspections"))


@app.route("/inspection/ppa/<int:ppa_id>")
@login_required
def inspection_ppa(ppa_id):
    u=current_user(); c=db()
    p=c.execute("SELECT * FROM ppa_establishments WHERE id=%s AND active=1",(ppa_id,)).fetchone()
    if not p:
        c.close(); return "PPA not found",404
    if u["role"]=="Supporting Staff":
        allowed=c.execute("SELECT 1 FROM inspection_assignments WHERE ppa_id=%s AND supporting_staff_id=%s AND status='Active'",(ppa_id,u["id"])).fetchone()
        if not allowed:
            c.close(); flash("This PPA is not assigned to you.","error"); return redirect(url_for("inspections"))
    elif u["role"] not in ("Administrator","Zonal Inspector"):
        c.close(); flash("Inspection access is not available to this role.","error"); return redirect(url_for("dashboard"))
    members=c.execute("SELECT * FROM corps_members WHERE ppa_id=%s AND status='Active' ORDER BY full_name",(ppa_id,)).fetchall()
    last_submission=c.execute("""SELECT ir.inspected_at,ir.attendance_status,u.full_name AS inspector_name
        FROM inspection_records ir JOIN users u ON u.id=ir.inspected_by
        WHERE ir.ppa_id=%s ORDER BY ir.id DESC LIMIT 1""",(ppa_id,)).fetchone()
    c.close()
    return render_template("inspection_ppa.html",ppa=p,members=members,last_submission=last_submission)

@app.route("/inspection/ppa/<int:ppa_id>/search")
@login_required
def inspection_search(ppa_id):
    u=current_user(); state=request.args.get("state_code","").strip()
    if u["role"]!="Supporting Staff": return {"error":"Only Supporting Staff may use this search."},403
    if not state: return {"found":False,"message":"Enter a State Code to search."}
    c=db()
    allowed=c.execute("SELECT 1 FROM inspection_assignments WHERE ppa_id=%s AND supporting_staff_id=%s AND status='Active'",(ppa_id,u["id"])).fetchone()
    if not allowed: c.close(); return {"error":"PPA is not assigned to you."},403
    row=c.execute("""SELECT cm.*,p.name AS ppa_name,p.lga,
        (SELECT ir.attendance_status FROM inspection_records ir WHERE ir.corps_member_id=cm.id ORDER BY ir.id DESC LIMIT 1) AS last_inspection_status,
        (SELECT ir.inspected_at FROM inspection_records ir WHERE ir.corps_member_id=cm.id ORDER BY ir.id DESC LIMIT 1) AS last_inspection_at,
        (SELECT ir.remarks FROM inspection_records ir WHERE ir.corps_member_id=cm.id ORDER BY ir.id DESC LIMIT 1) AS last_inspection_remarks,
        (SELECT ir.other_reason FROM inspection_records ir WHERE ir.corps_member_id=cm.id ORDER BY ir.id DESC LIMIT 1) AS last_inspection_other_reason
        FROM corps_members cm JOIN ppa_establishments p ON p.id=cm.ppa_id
        WHERE cm.ppa_id=%s AND UPPER(REPLACE(cm.state_code,' ',''))=UPPER(REPLACE(%s,' ','')) AND cm.status='Active'""",(ppa_id,state)).fetchone()
    c.close()
    if not row: return {"found":False,"message":"No active corps member with that State Code was found under this PPA."}
    return {"found":True,"member":dict(row)}

@app.route("/inspection/record",methods=["POST"])
@login_required
def inspection_record():
    u=current_user()
    if u["role"]!="Supporting Staff":
        flash("Only Supporting Staff may submit inspection attendance.","error"); return redirect(url_for("inspections"))
    ppa_id=request.form.get("ppa_id",type=int); member_id=request.form.get("corps_member_id",type=int); status=request.form.get("attendance_status","").strip(); other=request.form.get("other_reason","").strip(); remarks=request.form.get("remarks","").strip()
    allowed_status={"Present","Absent","Leave","Sick Leave","Maternity Leave","Others"}
    if status not in allowed_status:
        flash("Please select a valid inspection status.","error"); return redirect(url_for("inspection_ppa",ppa_id=ppa_id))
    if status=="Others" and not other:
        flash("Please specify the reason for Other status.","error"); return redirect(url_for("inspection_ppa",ppa_id=ppa_id))
    c=db(); allowed=c.execute("SELECT 1 FROM inspection_assignments WHERE ppa_id=%s AND supporting_staff_id=%s AND status='Active'",(ppa_id,u["id"])).fetchone(); member=c.execute("SELECT * FROM corps_members WHERE id=%s AND ppa_id=%s AND status='Active'",(member_id,ppa_id)).fetchone()
    if not allowed or not member:
        c.close(); flash("The selected corps member or PPA is not assigned to you.","error"); return redirect(url_for("inspections"))
    iid=next_new_inspection_id(); t=now(); query=1 if status=="Absent" else 0
    c.execute("""INSERT INTO inspection_records(inspection_id,corps_member_id,ppa_id,inspected_by,attendance_status,query_issued,other_reason,remarks,inspected_at)
                 VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)""",(iid,member_id,ppa_id,u["id"],status,query,other,remarks,t))
    findings=f"Attendance status: {status}. Query issued: {'Yes' if query else 'No'}." + (f" Reason: {other}." if other else "")
    c.execute("""INSERT INTO inspections(inspection_id,lga,ppa_employer,corps_member,subject,findings,recommendations,inspected_by,inspected_at,status)
                 VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",(iid,member["ppa_id"] and c.execute("SELECT lga FROM ppa_establishments WHERE id=%s",(ppa_id,)).fetchone()["lga"],member["ppa_id"] and c.execute("SELECT name FROM ppa_establishments WHERE id=%s",(ppa_id,)).fetchone()["name"],member["full_name"],"PPA Attendance Inspection",findings,remarks,u["full_name"],t,"Completed"))
    c.commit(); c.close(); flash(f"{iid} recorded successfully. Query issued: {'YES' if query else 'NO'}.","success"); return redirect(url_for("inspection_ppa",ppa_id=ppa_id))

@app.route("/reports")
@login_required
def reports():
    u=current_user(); c=db()
    if u["role"] == "LGI Officer": rows=c.execute("SELECT * FROM reports WHERE lga=%s ORDER BY id DESC",(u["lga"],)).fetchall()
    elif u["role"] in ("Administrator", "Zonal Inspector"): rows=c.execute("SELECT * FROM reports ORDER BY id DESC").fetchall()
    else: c.close(); flash("Reports are not available to this role.","error"); return redirect(url_for("dashboard"))
    c.close(); return render_template("reports.html",rows=rows,report_only=(u["role"]=="LGI Officer"))

@app.route("/report/new",methods=["GET","POST"])
@report_management_required
def new_report():
    if request.method=="POST":
        subject=request.form.get("subject","").strip(); body=request.form.get("report_body","").strip(); lga=request.form.get("lga","").strip()
        if not subject or not body or lga not in LGAS: flash("LGA, report subject and report content are required.","error"); return redirect(url_for("new_report"))
        t=now(); rid=next_report_id(); u=current_user(); c=db(); c.execute("INSERT INTO reports (report_id,lga,report_type,ppa_employer,corps_member,subject,report_body,issued_by,issued_at,status) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",(rid,lga,request.form.get("report_type","General LGA Report"),request.form.get("ppa_employer"),request.form.get("corps_member"),subject,body,u["full_name"],t,"Issued")); c.commit(); c.close(); flash(f"{rid} issued to {lga} LGI.","success"); return redirect(url_for("reports"))
    return render_template("report_form.html",lgas=LGAS)

@app.route("/download/<name>")
@login_required
@lgi_report_only
def download(name):
    return send_from_directory(UPLOADS,name,as_attachment=True)

@app.route("/admin/db-inspection")
@admin_required
def db_inspection():
    with db() as c:
        tables = c.execute("""
            SELECT tablename
            FROM pg_catalog.pg_tables
            WHERE schemaname='public'
            ORDER BY tablename
        """).fetchall()

        counts = {}
        for row in tables:
            table = row["tablename"]
            counts[table] = c.execute(
                f'SELECT COUNT(*) AS count FROM "{table}"'
            ).fetchone()["count"]

        database = c.execute(
            "SELECT current_database() AS database"
        ).fetchone()["database"]
        version = c.execute(
            "SELECT version() AS version"
        ).fetchone()["version"]

    rows = "".join(
        f"<tr><td>{name}</td><td>{count}</td></tr>"
        for name, count in counts.items()
    )
    html = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Database Inspection</title>
<style>body{{font-family:Arial,sans-serif;margin:40px}}table{{border-collapse:collapse}}
th,td{{border:1px solid #ccc;padding:8px 12px}}th{{background:#f3f3f3}}</style></head>
<body><h1>PostgreSQL Database Inspection</h1>
<p><b>Database:</b> {database}</p><p><b>Server:</b> {version}</p>
<table><tr><th>Table</th><th>Rows</th></tr>{rows}</table>
</body></html>"""
    return html

init_db()

if __name__=="__main__":
    app.run(debug=True)
